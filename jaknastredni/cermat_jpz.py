"""Importér agregovaných výsledků jednotné přijímací zkoušky CERMAT (JPZ),
nový formát 2024+, po školách/oborech.

Zdroj: `/files/files/JPZ/agregovana_data_skoly/PZ{rok}_kolo{kolo}_skolobory_{soubor}.xlsx`
na data.cermat.cz, roky 2024+, kola 1 a 2, soubor `vysledky`/`prihlasky`/`kapacity`.
Podrobný průzkum: docs/research/cermat.md, oddíly 6, 9, 10, 13.

Na rozdíl od starého formátu (2017–2023, tabulka `jpz_skupina`, dosud
neimplementováno) a od maturity je hlavička u nového formátu na **1. řádku**
listu (žádný sloučený titulek) a sloupce se mapují podle jména, ne pozice.
Jeden řádek = kombinace škola (IZO) × obor (KKOV) × zaměření oboru × forma
vzdělávání × délka studia × jazyk studia × ročník × rok × kolo — teprve tahle
rozšířená kombinace je v datech jednoznačná (viz `PRIMARY KEY` v schema.sql
a oddíl 13 průzkumu).

Tři soubory (`vysledky`, `prihlasky`, `kapacity`) popisují stejnou množinu
nabídek a spojují se přes `ID_SOF` (CERMAT UUID řádku/zaměření) — ověřeno,
100% shoda mezi soubory. `ID_SO` (UUID nabídky, sdílené mezi zaměřeními pod
jedním KKOV) je v `_prihlasky.xlsx`/`_kapacity.xlsx` za roky 2024–2025
přejmenované na `IS_SO` (patrně překlep na straně CERMAT) — `vysledky.xlsx`
má vždy `ID_SO`. `vysledky.xlsx` navíc obsahuje jako jediný ze tří souborů
kompletní sadu sloupců (kapacita, přihlášky, přijetí, skóre, percentily,
důvody nepřijetí) — `prihlasky`/`kapacita` jsou jeho podmnožinou; importují
se přesto všechny tři a spojují přes `ID_SOF`, jak žádá zadání, pro odolnost
vůči budoucím letům, kde by se sady sloupců mohly rozejít.

Použití:
    python -m jaknastredni.cermat_jpz --db data/jaknastredni.db --roky 2024-2026
"""
from __future__ import annotations

import argparse
import hashlib
import logging
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator

from . import db

log = logging.getLogger(__name__)

BASE_URL = ("https://data.cermat.cz/files/files/JPZ/agregovana_data_skoly/"
            "PZ{rok}_kolo{kolo}_skolobory_{soubor}.xlsx")
USER_AGENT = "jaknastredni/0.1 (+https://github.com/austy246/jaknastredni)"
SOUBORY = ("vysledky", "prihlasky", "kapacity")
KOLA = (1, 2)

Row = dict[str, Any]

# Sloupce společné identifikaci nabídky, stejné ve všech třech souborech.
_IDENT = ["ROK", "KOLO", "IZO", "REDIZO", "ROČNÍK", "KKOV",
          "ZAMĚŘENÍ OBORU", "FORMA VZDĚLÁVÁNÍ", "DÉLKA STUDIA", "JAZYK STUDIA"]

# Mapování název sloupce v hlavičce -> název sloupce v DB, po jednotlivých
# souborech. `_vysledky.xlsx` je nadmnožina `_kapacity.xlsx`/`_prihlasky.xlsx`
# (viz docstring modulu) - mapy se proto překrývají v KAPACITA/INDEX
# POPTÁVKY/PŘIHLÁŠKY sloupcích.
KAPACITA_COLS = {
    "KAPACITA": "kapacita",
}

PRIHLASKY_COLS = {
    "KAPACITA": "kapacita",
    "INDEX POPTÁVKY (PŘIHLÁŠKY / KAPACITA)": "index_poptavky",
    "PŘIHLÁŠKY CELKEM": "prihlasky_celkem",
    "PŘIHLÁŠKY - PRIORITA 1": "prihlasky_priorita_1",
    "PŘIHLÁŠKY - PRIORITA 2": "prihlasky_priorita_2",
    "PŘIHLÁŠKY - PRIORITA 3": "prihlasky_priorita_3",
    "PŘIHLÁŠKY - PRIORITA 4": "prihlasky_priorita_4",
    "PŘIHLÁŠKY - PRIORITA 5": "prihlasky_priorita_5",
}

VYSLEDKY_COLS = {
    **PRIHLASKY_COLS,
    "PŘIJATÍ": "prijati",
    "PŘIJATÍ - PRIORITA 1": "prijati_priorita_1",
    "PŘIJATÍ - PRIORITA 2": "prijati_priorita_2",
    "PŘIJATÍ - PRIORITA 3": "prijati_priorita_3",
    "PŘIJATÍ - PRIORITA 4": "prijati_priorita_4",
    "PŘIJATÍ - PRIORITA 5": "prijati_priorita_5",
    "ČJ+MA - KONALI": "konali_cjma",
    "ČJ - KONALI": "konali_cj",
    "MA - KONALI": "konali_ma",
    "ČJ+MA - % SKÓR - PRŮMĚR": "skor_prumer_cjma",
    "ČJ - % SKÓR - PRŮMĚR": "skor_prumer_cj",
    "MA - % SKÓR - PRŮMĚR": "skor_prumer_ma",
    "ČJ+MA - % SKÓR - MIN": "skor_min_cjma",
    "ČJ - % SKÓR - MIN": "skor_min_cj",
    "MA - % SKÓR - MIN": "skor_min_ma",
    "ČJ+MA - % SKÓR - MAX": "skor_max_cjma",
    "ČJ - % SKÓR - MAX": "skor_max_cj",
    "MA - % SKÓR - MAX": "skor_max_ma",
    "ČJ+MA - PERCENTIL - PRŮMĚR": "percentil_prumer_cjma",
    "ČJ - PERCENTIL - PRŮMĚR": "percentil_prumer_cj",
    "MA - PERCENTIL - PRŮMĚR": "percentil_prumer_ma",
    "ČJ+MA - PERCENTIL - MIN": "percentil_min_cjma",
    "ČJ - PERCENTIL - MIN": "percentil_min_cj",
    "MA - PERCENTIL - MIN": "percentil_min_ma",
    "ČJ+MA - PERCENTIL - MAX": "percentil_max_cjma",
    "ČJ - PERCENTIL - MAX": "percentil_max_cj",
    "MA - PERCENTIL - MAX": "percentil_max_ma",
    "ČJ+MA - KONALI (PŘIJATI)": "konali_prijati_cjma",
    "ČJ - KONALI (PŘIJATI)": "konali_prijati_cj",
    "MA - KONALI (PŘIJATI)": "konali_prijati_ma",
    "ČJ+MA - % SKÓR - PRŮMĚR (PŘIJATI)": "skor_prijati_prumer_cjma",
    "ČJ - % SKÓR - PRŮMĚR (PŘIJATI)": "skor_prijati_prumer_cj",
    "MA - % SKÓR - PRŮMĚR (PŘIJATI)": "skor_prijati_prumer_ma",
    "ČJ+MA - % SKÓR - MIN (PŘIJATI)": "skor_prijati_min_cjma",
    "ČJ - % SKÓR - MIN (PŘIJATI)": "skor_prijati_min_cj",
    "MA - % SKÓR - MIN (PŘIJATI)": "skor_prijati_min_ma",
    "ČJ+MA - % SKÓR - MAX (PŘIJATI)": "skor_prijati_max_cjma",
    "ČJ - % SKÓR - MAX (PŘIJATI)": "skor_prijati_max_cj",
    "MA - % SKÓR - MAX (PŘIJATI)": "skor_prijati_max_ma",
    "ČJ+MA - PERCENTIL - PRŮMĚR (PŘIJATI)": "percentil_prijati_prumer_cjma",
    "ČJ - PERCENTIL - PRŮMĚR (PŘIJATI)": "percentil_prijati_prumer_cj",
    "MA - PERCENTIL - PRŮMĚR (PŘIJATI)": "percentil_prijati_prumer_ma",
    "ČJ+MA - PERCENTIL - MIN (PŘIJATI)": "percentil_prijati_min_cjma",
    "ČJ - PERCENTIL - MIN (PŘIJATI)": "percentil_prijati_min_cj",
    "MA - PERCENTIL - MIN (PŘIJATI)": "percentil_prijati_min_ma",
    "ČJ+MA - PERCENTIL - MAX (PŘIJATI)": "percentil_prijati_max_cjma",
    "ČJ - PERCENTIL - MAX (PŘIJATI)": "percentil_prijati_max_cj",
    "MA - PERCENTIL - MAX (PŘIJATI)": "percentil_prijati_max_ma",
    "NEPŘIJATI - PŘIJAT NA VYŠŠÍ PRIORITU": "neprijati_vyssi_priorita",
    "NEPŘIJATI - NEDOSTATEČNÁ KAPACITA": "neprijati_nedostatecna_kapacita",
    "NEPŘIJATI - NESPLNĚNÍ PODMÍNEK": "neprijati_nesplneni_podminek",
    "NEPŘIJATI - VZDAL SE PŘIJETÍ": "neprijati_vzdal_se",
}

_COLMAPS = {"kapacity": KAPACITA_COLS, "prihlasky": PRIHLASKY_COLS, "vysledky": VYSLEDKY_COLS}

# Všechny metrické sloupce (sjednocení výše) - použije se pro doplnění None
# u řádků, které se nepodařilo spojit se všemi třemi soubory.
_ALL_METRIC_COLS = sorted(set(VYSLEDKY_COLS.values()))

DB_COLS = (
    ["izo", "kod_kkov", "rocnik", "rok", "kolo", "zamereni_oboru", "forma_vzdelavani",
     "delka_studia", "jazyk_studia", "redizo", "id_so", "id_sof"]
    + _ALL_METRIC_COLS
)


# --------------------------------------------------------------------------- stažení

def download(rok: int, kolo: int, soubor: str, raw_dir: Path = Path("data/raw/cermat"),
             timeout: int = 120) -> Path | None:
    """Stáhne jeden ze tří souborů (vysledky/prihlasky/kapacity) pro rok+kolo.

    Vrací None (a zaloguje varování), pokud soubor na CERMATu ještě/už
    neexistuje (HTTP 404) - typicky 2. kolo posledního ročníku, dokud
    neskončí přijímací řízení. Jiné HTTP chyby se propagují.
    """
    import requests  # lokální import, ať jdou testy bez sítě

    url = BASE_URL.format(rok=rok, kolo=kolo, soubor=soubor)
    log.info("Stahuji %s", url)
    resp = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=timeout)
    if resp.status_code == 404:
        log.warning("%s: soubor zatím neexistuje (HTTP 404), přeskakuji", url)
        return None
    resp.raise_for_status()
    raw_dir.mkdir(parents=True, exist_ok=True)
    path = raw_dir / f"PZ{rok}_kolo{kolo}_skolobory_{soubor}.xlsx"
    path.write_bytes(resp.content)
    log.info("Uloženo %s (%d B)", path, len(resp.content))
    return path


# --------------------------------------------------------------------------- parsování

def _redizo(v: Any) -> str:
    """REDIZO má 9 číslic; v souborech přichází občas jako číslo, občas jako text."""
    if isinstance(v, float):
        v = int(v)
    return str(v).strip().zfill(9)


def _izo(v: Any) -> str:
    """IZO má v JPZ souborech textový prefix 'izo_' - odstranit (viz README,
    docs/datovy-model.md "Spojování zdrojů")."""
    s = str(v).strip()
    if s.startswith("izo_"):
        s = s[len("izo_"):]
    return s.zfill(9)


def _delka(v: Any) -> str:
    """DÉLKA STUDIA přichází v kapacity.xlsx jako text ('4.0'), ve vysledky.xlsx
    jako číslo (4) - normalizovat na jednotný textový tvar bez desetiny."""
    return str(int(round(float(v))))


def _num(v: Any) -> float | int | None:
    """Stejná konvence jako v cermat_mz: '-'/prázdno = NULL."""
    if v is None or v == "" or v == "-":
        return None
    return v


def _sheet_header_and_rows(path: Path):
    """Vrátí (workbook, header, řádky) - list, který není 'vysvetlivky'.
    Na rozdíl od starého JPZ formátu a maturity je hlavička nového formátu
    JPZ na 1. řádku (žádný sloučený titulek), viz cermat.md oddíl 6/13."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = next(s for s in wb.sheetnames if s != "vysvetlivky")
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(min_row=1, values_only=True)
    header = list(next(rows_iter))
    return wb, header, rows_iter


def parse(paths: dict[str, Path]) -> Iterator[Row]:
    """Spojí vysledky/prihlasky/kapacity přes ID_SOF (CERMAT UUID řádku -
    ověřeno jako spolehlivý klíč, viz docstring modulu) a vrátí sloučené řádky.

    `paths` mapuje název souboru ('vysledky'/'prihlasky'/'kapacity') na cestu;
    chybějící klíče (soubor nevyšel stáhnout, viz `download`) se přeskočí.
    Pořadí zpracování kapacity -> prihlasky -> vysledky zajišťuje, že
    nejúplnější zdroj (vysledky) má u překrývajících se sloupců poslední/
    rozhodující slovo.
    """
    merged: dict[str, Row] = {}

    for soubor in ("kapacity", "prihlasky", "vysledky"):
        path = paths.get(soubor)
        if path is None:
            continue
        wb, header, rows_iter = _sheet_header_and_rows(path)
        try:
            idx = {h: i for i, h in enumerate(header)}
            missing_ident = [c for c in _IDENT if c not in idx]
            if missing_ident or "ID_SOF" not in idx:
                raise ValueError(f"{path}: chybí očekávané sloupce {missing_ident or ['ID_SOF']}")
            id_so_col = "ID_SO" if "ID_SO" in idx else ("IS_SO" if "IS_SO" in idx else None)
            colmap = _COLMAPS[soubor]

            n_rows = 0
            for row in rows_iter:
                n_rows += 1
                id_sof = row[idx["ID_SOF"]]
                if id_sof is None:
                    continue
                entry = merged.setdefault(id_sof, {"id_sof": id_sof, "id_so": None})
                entry["izo"] = _izo(row[idx["IZO"]])
                entry["redizo"] = _redizo(row[idx["REDIZO"]])
                entry["kod_kkov"] = row[idx["KKOV"]]
                entry["rocnik"] = int(row[idx["ROČNÍK"]])
                entry["rok"] = int(row[idx["ROK"]])
                entry["kolo"] = int(row[idx["KOLO"]])
                entry["zamereni_oboru"] = row[idx["ZAMĚŘENÍ OBORU"]] or ""
                entry["forma_vzdelavani"] = row[idx["FORMA VZDĚLÁVÁNÍ"]]
                entry["delka_studia"] = _delka(row[idx["DÉLKA STUDIA"]])
                entry["jazyk_studia"] = row[idx["JAZYK STUDIA"]]
                if id_so_col is not None and row[idx[id_so_col]] is not None:
                    entry["id_so"] = row[idx[id_so_col]]
                for hdr, dbcol in colmap.items():
                    if hdr in idx:
                        entry[dbcol] = _num(row[idx[hdr]])
            if n_rows == 0:
                log.warning("%s: list je prázdný (0 řádků dat), přeskakuji", path)
        finally:
            wb.close()

    for entry in merged.values():
        for dbcol in _ALL_METRIC_COLS:
            entry.setdefault(dbcol, None)
        yield entry


# --------------------------------------------------------------------------- import

def import_rows(conn: sqlite3.Connection, rows: Iterable[Row], *, url: str | None = None,
                 soubor: Path | None = None, rok: int | None = None, kolo: int | None = None,
                 jen_izo: set[str] | None = None) -> dict[str, int]:
    """Vloží řádky do tabulky prijimaci_rizeni (INSERT OR REPLACE, idempotentní)."""
    rows = list(rows)
    if jen_izo is not None:
        rows = [r for r in rows if r["izo"] in jen_izo]
    sql = f"INSERT OR REPLACE INTO prijimaci_rizeni ({', '.join(DB_COLS)}) VALUES ({', '.join('?' for _ in DB_COLS)})"
    with conn:
        if rows:
            conn.executemany(sql, [tuple(r.get(c) for c in DB_COLS) for r in rows])
        sha = hashlib.sha256(soubor.read_bytes()).hexdigest() if soubor and soubor.exists() else None
        conn.execute(
            "INSERT INTO import_run (zdroj, url, soubor, sha256, datum_vystupu, stazeno_utc, pocet_zaznamu, poznamka)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            ("cermat_jpz", url, str(soubor) if soubor else None, sha,
             str(rok) if rok else None, datetime.now(timezone.utc).isoformat(timespec="seconds"),
             len(rows), f"kolo {kolo}" if kolo else None),
        )
    return {"prijimaci_rizeni": len(rows)}


# --------------------------------------------------------------------------- CLI

def _parse_roky(spec: str) -> list[int]:
    m = re.fullmatch(r"(\d{4})-(\d{4})", spec)
    if m:
        return list(range(int(m.group(1)), int(m.group(2)) + 1))
    return [int(x) for x in spec.split(",")]


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Import agregovaných výsledků JPZ CERMAT (nový formát 2024+) do SQLite.")
    p.add_argument("--db", default="data/jaknastredni.db", help="cesta k SQLite databázi")
    p.add_argument("--roky", default=str(datetime.now().year), help="rok nebo rozsah, např. 2024-2026")
    p.add_argument("--raw-dir", type=Path, default=Path("data/raw/cermat"), help="kam ukládat stažené soubory")
    p.add_argument("--jen-praha", action="store_true",
                    help="importovat jen IZO přítomná v tabulce skola")
    p.add_argument("-v", "--verbose", action="store_true")
    args = p.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(levelname)s %(message)s")

    conn = db.connect(args.db)
    jen_izo = None
    if args.jen_praha:
        jen_izo = {r[0] for r in conn.execute("SELECT izo FROM skola")}

    total = 0
    for rok in _parse_roky(args.roky):
        for kolo in KOLA:
            paths: dict[str, Path] = {}
            for soubor in SOUBORY:
                path = download(rok, kolo, soubor, args.raw_dir)
                if path is not None:
                    paths[soubor] = path
            if "vysledky" not in paths:
                log.warning("PZ%s kolo %d: chybí vysledky.xlsx, kolo přeskočeno", rok, kolo)
                continue
            rows = list(parse(paths))
            if not rows:
                log.warning("PZ%s kolo %d: po naparsování 0 řádků, přeskočeno", rok, kolo)
                continue
            stats = import_rows(conn, rows, url=BASE_URL.format(rok=rok, kolo=kolo, soubor="vysledky"),
                                 soubor=paths["vysledky"], rok=rok, kolo=kolo, jen_izo=jen_izo)
            log.info("PZ%s kolo %d: %s (soubory: %s)", rok, kolo, stats, sorted(paths))
            total += stats["prijimaci_rizeni"]

    n = conn.execute("SELECT COUNT(*) FROM prijimaci_rizeni").fetchone()[0]
    log.info("Import hotov, %d řádků naimportováno/aktualizováno v tomto běhu, tabulka prijimaci_rizeni má celkem %d řádků", total, n)
    return 0


if __name__ == "__main__":
    sys.exit(main())
