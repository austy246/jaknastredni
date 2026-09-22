"""Importér souborů uchazečů CERMAT (JPZ 2024+) → tabulka `prijimacky_pasmo`.

Zdroj: `/files/JPZ-polozkova-data/{rok}/Uchazeci/PZ{rok}_kolo{kolo}_uchazeci_prihlasky_vysledky.xlsx`
na data.cermat.cz, roky 2024+, kola 1 a 2. Průzkum: docs/research/cermat.md,
oddíl 15.

**Proč tenhle zdroj.** Agregovaná data po školách (`prijimaci_rizeni`) dávají
jako jedinou informaci o náročnosti *minimální* % skór přijatého uchazeče. To
je ocasová hodnota — často jeden člověk, který se dostal na body za prospěch
nebo v rozřazení při shodě. Odhad šance postavený na ní je systematicky
optimistický. Soubory uchazečů dovolují spočítat **skutečný podíl přijatých
v každém bodovém pásmu**, tedy tu veličinu, kterou průvodce opravdu chce.

**Co soubor obsahuje a co ne.** Jeden řádek = jeden (anonymizovaný) uchazeč:
% skór ČJ+MA / ČJ / MA a až **pět přihlášek** (`ss1..ss5`) s REDIZO, KKOV,
zřizovatelem, formou, příznakem přijetí a důvodem nepřijetí. Neobsahuje
**žádný údaj o základní škole ani o známkách** — odhadnout šanci „podle
vysvědčení" z veřejných dat tedy nejde, a tenhle soubor je nejblíž, co
veřejně existuje.

**Co se do tabulky nepočítá.** Uchazeč nepřijatý s důvodem
`prijat_na_vyssi_prioritu` nebyl posouzen věcně — škola ho z pořadí vyřadila
proto, že se dostal jinam výš. Do míry přijetí patří jen přihlášky, které
věcné posouzení prošly (přijat, `pro_nedostacujici_kapacitu`,
`pro_nesplneni_podminek`); vyšší priorita a `vzdal_se` se ukládají zvlášť,
ale do jmenovatele nevstupují (viz `mira_prijeti` a `docs/pruvodce-ux.md`).

Zjištění ověřená na stažených souborech (22. 9. 2026):
- Hlavička je na **1. řádku**, list se jmenuje `Sheet 1`, druhý list
  `legenda` je datový slovník. Sloupců je 40 — kromě **PZ2026 kolo 2, který
  má navíc `rocnik`** (41). Mapuje se proto podle jména, ne pozice.
- `ss*_prijat`: `1` = přijat, `2` = nepřijat. `ss*_forma`: `den`, `dal`,
  `vec`, `dist`, `komb`, `den2`. `ss*_zkraceno`: `1` = ano, `2` = ne.
- Importují se jen přihlášky na **denní nezkrácené** studium (`forma == 'den'`
  a `zkraceno != '1'`) — průvodce jiné neřeší.
- V roce 2026 (1. kolo) je 156 210 uchazečů, z toho **37 004 bez % skóru**;
  jsou to skoro výhradně obory s výučním listem (H/E), které JPZ nekonají.
  Ukládají se do pásma `-1`, aby i u nich šlo spočítat míru přijetí.
- Dataset obsahuje „pouze platné přihlášky" ke dni uzávěrky (list `legenda`).

Použití:
    python -m jaknastredni.cermat_uchazeci --db data/jaknastredni.db --roky 2024-2026
"""
from __future__ import annotations

import argparse
import hashlib
import logging
import re
import sqlite3
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator

from . import db

log = logging.getLogger(__name__)

BASE_URL = ("https://data.cermat.cz/files/JPZ-polozkova-data/{rok}/Uchazeci/"
            "PZ{rok}_kolo{kolo}_uchazeci_prihlasky_vysledky.xlsx")
USER_AGENT = "jaknastredni/0.1 (+https://github.com/austy246/jaknastredni)"
KOLA = (1, 2)

# Šířka bodového pásma na škále 0–200. Pět bodů je jemnější, než co průvodce
# zobrazuje, ale drží dost dat na to, aby šlo pásma při dotazu slučovat.
SIRKA_PASMA = 5
PASMO_BEZ_JPZ = -1

# Důvody nepřijetí, jak je uvádí sloupec ss*_duvod_neprijeti.
DUVOD_KAPACITA = "pro_nedostacujici_kapacitu"
DUVOD_PODMINKY = "pro_nesplneni_podminek"
DUVOD_VYSSI = "prijat_na_vyssi_prioritu"
DUVOD_VZDAL = "vzdal_se"

DB_COLS = ["redizo", "kod_kkov", "rok", "kolo", "pasmo_od", "prihlasek", "prijato",
           "nedostatecna_kapacita", "nesplneni_podminek", "vyssi_priorita", "vzdal_se"]

Row = dict[str, Any]


# --------------------------------------------------------------------------- stažení

def download(rok: int, kolo: int, raw_dir: Path = Path("data/raw/cermat"),
             timeout: int = 300) -> Path | None:
    """Stáhne soubor uchazečů pro rok+kolo (~16 MB u 1. kola, ~1 MB u 2.).

    Vrací None, pokud soubor ještě neexistuje (HTTP 404) — typicky 2. kolo
    probíhajícího ročníku. Stejná konvence jako `cermat_jpz.download`.
    """
    import requests  # lokální import, ať jdou testy bez sítě

    url = BASE_URL.format(rok=rok, kolo=kolo)
    log.info("Stahuji %s", url)
    resp = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=timeout)
    if resp.status_code == 404:
        log.warning("%s: soubor zatím neexistuje (HTTP 404), přeskakuji", url)
        return None
    resp.raise_for_status()
    raw_dir.mkdir(parents=True, exist_ok=True)
    path = raw_dir / f"PZ{rok}_kolo{kolo}_uchazeci_prihlasky_vysledky.xlsx"
    path.write_bytes(resp.content)
    log.info("Uloženo %s (%d B)", path, len(resp.content))
    return path


# --------------------------------------------------------------------------- parsování

def _redizo(v: Any) -> str:
    """REDIZO má 9 číslic; přichází jako text i jako číslo."""
    if isinstance(v, float):
        v = int(v)
    return str(v).strip().zfill(9)


def _pasmo(skor: Any) -> int:
    """Dolní mez pásma pro % skór ČJ+MA; PASMO_BEZ_JPZ, když uchazeč nekonal."""
    if skor is None or skor == "":
        return PASMO_BEZ_JPZ
    try:
        hodnota = float(skor)
    except (TypeError, ValueError):
        return PASMO_BEZ_JPZ
    return int(hodnota // SIRKA_PASMA) * SIRKA_PASMA


def parse(path: Path, rok: int, kolo: int) -> Iterator[Row]:
    """Agreguje uchazeče ze souboru na (REDIZO × KKOV × pásmo) s počty.

    Každý uchazeč přispěje až pěti přihláškami (priority 1–5). Rok a kolo se
    berou z parametrů, ne ze sloupců souboru — sloupce sice existují, ale
    název souboru je autoritativní a konzistentní s ostatními importéry.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[next(s for s in wb.sheetnames if s != "legenda")]
        radky = ws.iter_rows(min_row=1, values_only=True)
        hlavicka = [str(h) if h is not None else "" for h in next(radky)]
        idx = {h: i for i, h in enumerate(hlavicka)}
        chybi = [c for c in ("c_m_procentni_skor", "ss1_redizo", "ss1_kkov") if c not in idx]
        if chybi:
            raise ValueError(f"{path.name}: v hlavičce chybí sloupce {chybi}")

        kose: dict[tuple[str, str, int], dict[str, int]] = defaultdict(
            lambda: dict.fromkeys(
                ("prihlasek", "prijato", "nedostatecna_kapacita",
                 "nesplneni_podminek", "vyssi_priorita", "vzdal_se"), 0)
        )
        for radek in radky:
            pasmo = _pasmo(radek[idx["c_m_procentni_skor"]])
            for priorita in range(1, 6):
                redizo_raw = radek[idx.get(f"ss{priorita}_redizo", -1)] if f"ss{priorita}_redizo" in idx else None
                if not redizo_raw:
                    continue
                if str(radek[idx[f"ss{priorita}_forma"]] or "").strip() != "den":
                    continue        # dálkové/večerní/zkrácené průvodce neřeší
                if str(radek[idx[f"ss{priorita}_zkraceno"]] or "").strip() == "1":
                    continue
                kkov = str(radek[idx[f"ss{priorita}_kkov"]] or "").strip()
                if not kkov:
                    continue
                kos = kose[(_redizo(redizo_raw), kkov, pasmo)]
                kos["prihlasek"] += 1
                if str(radek[idx[f"ss{priorita}_prijat"]] or "").strip() == "1":
                    kos["prijato"] += 1
                else:
                    duvod = str(radek[idx[f"ss{priorita}_duvod_neprijeti"]] or "").strip()
                    if duvod == DUVOD_KAPACITA:
                        kos["nedostatecna_kapacita"] += 1
                    elif duvod == DUVOD_PODMINKY:
                        kos["nesplneni_podminek"] += 1
                    elif duvod == DUVOD_VYSSI:
                        kos["vyssi_priorita"] += 1
                    elif duvod == DUVOD_VZDAL:
                        kos["vzdal_se"] += 1
    finally:
        wb.close()

    for (redizo, kkov, pasmo), kos in kose.items():
        yield {"redizo": redizo, "kod_kkov": kkov, "rok": rok, "kolo": kolo,
               "pasmo_od": pasmo, **kos}


# --------------------------------------------------------------------------- import

def import_rows(conn: sqlite3.Connection, rows: Iterable[Row], *, url: str | None = None,
                soubor: Path | None = None, rok: int | None = None, kolo: int | None = None,
                jen_redizo: set[str] | None = None) -> dict[str, int]:
    """Vloží agregované řádky do `prijimacky_pasmo` (INSERT OR REPLACE)."""
    rows = list(rows)
    if jen_redizo is not None:
        rows = [r for r in rows if r["redizo"] in jen_redizo]
    sql = (f"INSERT OR REPLACE INTO prijimacky_pasmo ({', '.join(DB_COLS)})"
           f" VALUES ({', '.join('?' for _ in DB_COLS)})")
    with conn:
        if rows:
            conn.executemany(sql, [tuple(r[c] for c in DB_COLS) for r in rows])
        sha = hashlib.sha256(soubor.read_bytes()).hexdigest() if soubor and soubor.exists() else None
        conn.execute(
            "INSERT INTO import_run (zdroj, url, soubor, sha256, datum_vystupu, stazeno_utc,"
            " pocet_zaznamu, poznamka) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            ("cermat_uchazeci", url, str(soubor) if soubor else None, sha,
             str(rok) if rok else None,
             datetime.now(timezone.utc).isoformat(timespec="seconds"),
             len(rows), f"kolo {kolo}" if kolo else None),
        )
    return {"prijimacky_pasmo": len(rows)}


# --------------------------------------------------------------------------- dotaz

def mira_prijeti(conn: sqlite3.Connection, redizo: str, kod_kkov: str, skor: float | None,
                 *, rozsah: int = 10, roky: Iterable[int] | None = None,
                 ) -> tuple[float, int] | None:
    """Podíl přijatých mezi uchazeči s podobným skórem — a velikost vzorku.

    `skor` je % skór ČJ+MA (0–200); None znamená „obor bez JPZ" a sáhne se do
    pásma `PASMO_BEZ_JPZ`. `rozsah` je poloviční šířka okna kolem skóru
    (výchozích ±10 bodů); sčítá se přes všechny uvedené roky i kola.

    Jmenovatel jsou jen věcně posouzené přihlášky — přijatí plus nepřijatí
    kvůli kapacitě nebo nesplnění podmínek. Kdo se dostal na vyšší prioritu
    nebo se vzdal, se nepočítá ani do čitatele, ani do jmenovatele.

    Vrací None, když v okně není ani jedna posouzená přihláška; volající si
    rozhodne, jestli okno rozšíří, nebo sáhne po odhadu z hranice přijetí.
    """
    podminka_roky = ""
    parametry: list[Any] = [redizo, kod_kkov]
    if roky is not None:
        roky = list(roky)
        podminka_roky = f" AND rok IN ({', '.join('?' for _ in roky)})"
        parametry += roky
    if skor is None:
        podminka_pasmo = " AND pasmo_od = ?"
        parametry.append(PASMO_BEZ_JPZ)
    else:
        podminka_pasmo = " AND pasmo_od >= ? AND pasmo_od <= ?"
        parametry += [max(0, int(skor) - rozsah), int(skor) + rozsah]

    radek = conn.execute(
        "SELECT SUM(prijato) AS prijato,"
        "       SUM(prijato + nedostatecna_kapacita + nesplneni_podminek) AS posouzeno"
        "  FROM prijimacky_pasmo"
        " WHERE redizo = ? AND kod_kkov = ?" + podminka_roky + podminka_pasmo,
        parametry,
    ).fetchone()
    posouzeno = radek[1] or 0
    if not posouzeno:
        return None
    return (radek[0] or 0) / posouzeno, posouzeno


# --------------------------------------------------------------------------- CLI

def _parse_roky(spec: str) -> list[int]:
    m = re.fullmatch(r"(\d{4})-(\d{4})", spec)
    if m:
        return list(range(int(m.group(1)), int(m.group(2)) + 1))
    return [int(x) for x in spec.split(",")]


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        description="Import souborů uchazečů CERMAT (JPZ 2024+) do SQLite.")
    p.add_argument("--db", default="data/jaknastredni.db", help="cesta k SQLite databázi")
    p.add_argument("--roky", default="2024-2026", help="rok nebo rozsah, např. 2024-2026")
    p.add_argument("--raw-dir", type=Path, default=Path("data/raw/cermat"),
                   help="kam ukládat stažené soubory")
    p.add_argument("--jen-praha", action="store_true",
                   help="importovat jen REDIZO přítomná v tabulce organizace")
    p.add_argument("-v", "--verbose", action="store_true")
    args = p.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(levelname)s %(message)s")

    conn = db.connect(args.db)
    jen_redizo = None
    if args.jen_praha:
        jen_redizo = {r[0] for r in conn.execute("SELECT redizo FROM organizace")}

    celkem = 0
    for rok in _parse_roky(args.roky):
        for kolo in KOLA:
            path = download(rok, kolo, args.raw_dir)
            if path is None:
                continue
            rows = list(parse(path, rok, kolo))
            if not rows:
                log.warning("PZ%s kolo %d: 0 řádků, přeskočeno", rok, kolo)
                continue
            stats = import_rows(conn, rows, url=BASE_URL.format(rok=rok, kolo=kolo),
                                soubor=path, rok=rok, kolo=kolo, jen_redizo=jen_redizo)
            log.info("PZ%s kolo %d: %s", rok, kolo, stats)
            celkem += stats["prijimacky_pasmo"]
    log.info("Hotovo, celkem %d řádků v prijimacky_pasmo.", celkem)
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
