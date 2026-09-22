from pathlib import Path

import openpyxl
import pytest

from jaknastredni import cermat_jpz, db

IDENT_HEADER = ["ID_SOF", "ID_SO", "ROK", "KOLO", "IZO", "REDIZO", "NÁZEV ŠKOLY",
                "ROČNÍK", "KKOV", "ZAMĚŘENÍ OBORU", "FORMA VZDĚLÁVÁNÍ",
                "DÉLKA STUDIA", "JAZYK STUDIA"]
KAPACITA_EXTRA = ["KAPACITA"]
PRIHLASKY_EXTRA = KAPACITA_EXTRA + [
    "INDEX POPTÁVKY (PŘIHLÁŠKY / KAPACITA)", "PŘIHLÁŠKY CELKEM",
    "PŘIHLÁŠKY - PRIORITA 1", "PŘIHLÁŠKY - PRIORITA 2",
    "PŘIHLÁŠKY - PRIORITA 3", "PŘIHLÁŠKY - PRIORITA 4", "PŘIHLÁŠKY - PRIORITA 5",
]
VYSLEDKY_EXTRA = PRIHLASKY_EXTRA + [
    "PŘIJATÍ", "PŘIJATÍ - PRIORITA 1", "PŘIJATÍ - PRIORITA 2", "PŘIJATÍ - PRIORITA 3",
    "PŘIJATÍ - PRIORITA 4", "PŘIJATÍ - PRIORITA 5",
    "ČJ+MA - KONALI", "ČJ - KONALI", "MA - KONALI",
    "ČJ+MA - % SKÓR - PRŮMĚR", "ČJ - % SKÓR - PRŮMĚR", "MA - % SKÓR - PRŮMĚR",
    "NEPŘIJATI - PŘIJAT NA VYŠŠÍ PRIORITU", "NEPŘIJATI - NEDOSTATEČNÁ KAPACITA",
    "NEPŘIJATI - NESPLNĚNÍ PODMÍNEK", "NEPŘIJATI - VZDAL SE PŘIJETÍ",
]


def _row(id_sof, id_so, rok, kolo, izo, redizo, nazev, rocnik, kkov, zamereni, forma, delka,
         jazyk, extra: dict, header_extra: list[str], id_so_col="ID_SO"):
    base = {
        "ID_SOF": id_sof, id_so_col: id_so, "ROK": rok, "KOLO": kolo, "IZO": izo,
        "REDIZO": redizo, "NÁZEV ŠKOLY": nazev, "ROČNÍK": rocnik, "KKOV": kkov,
        "ZAMĚŘENÍ OBORU": zamereni, "FORMA VZDĚLÁVÁNÍ": forma, "DÉLKA STUDIA": delka,
        "JAZYK STUDIA": jazyk,
    }
    base.update(extra)
    header = [h if h != "ID_SO" else id_so_col for h in IDENT_HEADER] + header_extra
    return [base.get(h) for h in header], header


def _build(path: Path, rows_spec: list[dict], header_extra: list[str], id_so_col="ID_SO"):
    wb = openpyxl.Workbook()
    ws = wb.active
    header = [h if h != "ID_SO" else id_so_col for h in IDENT_HEADER] + header_extra
    ws.append(header)  # hlavička je na 1. řádku (žádný sloučený titulek), viz cermat.md odd. 6/13
    for spec in rows_spec:
        row, _ = _row(header_extra=header_extra, id_so_col=id_so_col, **spec)
        ws.append(row)
    wb.save(path)


# Referenční obor: Obchodní akademie Heroldovy sady, REDIZO 600006573, KKOV 63-41-M/02
REF_IZO = "izo_000638510"
REF_REDIZO = 600006573
REF_KKOV = "63-41-M/02"


def _vysledky_row(id_sof, id_so, rok, kolo, izo=REF_IZO, redizo=REF_REDIZO,
                   nazev="Obchodní akademie", rocnik=9, kkov=REF_KKOV, zamereni="",
                   forma="den", delka=4, jazyk="Český", kapacita=30, prihlasky=282,
                   prijati=30, id_so_col="ID_SO"):
    extra = {
        "KAPACITA": kapacita,
        "INDEX POPTÁVKY (PŘIHLÁŠKY / KAPACITA)": round(prihlasky / kapacita, 2),
        "PŘIHLÁŠKY CELKEM": prihlasky,
        "PŘIHLÁŠKY - PRIORITA 1": prihlasky - 4, "PŘIHLÁŠKY - PRIORITA 2": 3,
        "PŘIHLÁŠKY - PRIORITA 3": 1, "PŘIHLÁŠKY - PRIORITA 4": 0, "PŘIHLÁŠKY - PRIORITA 5": 0,
        "PŘIJATÍ": prijati, "PŘIJATÍ - PRIORITA 1": prijati - 2, "PŘIJATÍ - PRIORITA 2": 2,
        "PŘIJATÍ - PRIORITA 3": 0, "PŘIJATÍ - PRIORITA 4": 0, "PŘIJATÍ - PRIORITA 5": 0,
        "ČJ+MA - KONALI": prihlasky - 10, "ČJ - KONALI": prihlasky - 10, "MA - KONALI": prihlasky - 10,
        "ČJ+MA - % SKÓR - PRŮMĚR": 61.7, "ČJ - % SKÓR - PRŮMĚR": 63.9, "MA - % SKÓR - PRŮMĚR": 58.0,
        "NEPŘIJATI - PŘIJAT NA VYŠŠÍ PRIORITU": prihlasky - prijati - 2,
        "NEPŘIJATI - NEDOSTATEČNÁ KAPACITA": 0, "NEPŘIJATI - NESPLNĚNÍ PODMÍNEK": 2,
        "NEPŘIJATI - VZDAL SE PŘIJETÍ": 0,
    }
    return dict(id_sof=id_sof, id_so=id_so, rok=rok, kolo=kolo, izo=izo, redizo=redizo,
                nazev=nazev, rocnik=rocnik, kkov=kkov, zamereni=zamereni, forma=forma,
                delka=delka, jazyk=jazyk, extra=extra)


@pytest.fixture
def rok_kolo(tmp_path):
    return 2026, 1


@pytest.fixture
def new_style_files(tmp_path, rok_kolo):
    """2026-styl: ID_SO pojmenovaný stejně ve všech třech souborech."""
    rok, kolo = rok_kolo
    rows = [
        _vysledky_row("sof-1", "so-1", rok, kolo),
        # druhá škola/obor, pro ověření počtu řádků
        _vysledky_row("sof-2", "so-2", rok, kolo, izo="izo_000012345", redizo=600012345,
                      nazev="Gymnázium Testovací", kkov="79-41-K/41", kapacita=60, prihlasky=120, prijati=60),
        # stejné IZO+KKOV+ROČNÍK+ROK+KOLO jako sof-1, ale jiné zaměření oboru
        # (ověřuje rozšířený primární klíč, viz schema.sql / cermat.md odd. 13)
        _vysledky_row("sof-3", "so-1", rok, kolo, zamereni="Ekonomika a finance", kapacita=15, prihlasky=40, prijati=15),
    ]
    vysledky = tmp_path / "vysledky.xlsx"
    prihlasky = tmp_path / "prihlasky.xlsx"
    kapacity = tmp_path / "kapacity.xlsx"
    _build(vysledky, rows, VYSLEDKY_EXTRA)
    _build(prihlasky, rows, PRIHLASKY_EXTRA)
    _build(kapacity, rows, KAPACITA_EXTRA)
    return {"vysledky": vysledky, "prihlasky": prihlasky, "kapacity": kapacity}


@pytest.fixture
def old_style_prihlasky_kapacity(tmp_path, rok_kolo):
    """2024/2025-styl: prihlasky/kapacity mají sloupec IS_SO místo ID_SO
    (viz docs/research/cermat.md oddíl 13) - join musí fungovat i tak."""
    rok, kolo = rok_kolo
    rows = [_vysledky_row("sof-1", "so-1", rok, kolo)]
    prihlasky = tmp_path / "prihlasky_old.xlsx"
    kapacity = tmp_path / "kapacity_old.xlsx"
    _build(prihlasky, rows, PRIHLASKY_EXTRA, id_so_col="IS_SO")
    _build(kapacity, rows, KAPACITA_EXTRA, id_so_col="IS_SO")
    return {"prihlasky": prihlasky, "kapacity": kapacity}


def test_parse_row_count_and_izo_prefix_stripped(new_style_files):
    rows = list(cermat_jpz.parse(new_style_files))
    assert len(rows) == 3
    assert all(not r["izo"].startswith("izo_") for r in rows)
    assert {r["izo"] for r in rows} == {"000638510", "000012345"}


def test_extended_primary_key_disambiguates_zamereni(new_style_files):
    rows = list(cermat_jpz.parse(new_style_files))
    ref_rows = [r for r in rows if r["izo"] == "000638510"]
    assert len(ref_rows) == 2  # sof-1 a sof-3, stejné izo+kkov+rocnik+rok+kolo
    assert {r["zamereni_oboru"] for r in ref_rows} == {"", "Ekonomika a finance"}
    keys = {(r["izo"], r["kod_kkov"], r["rocnik"], r["rok"], r["kolo"],
             r["zamereni_oboru"], r["forma_vzdelavani"], r["delka_studia"], r["jazyk_studia"])
            for r in ref_rows}
    assert len(keys) == 2  # klíč je jednoznačný, i když (izo,kkov,rocnik,rok,kolo) sám o sobě není


def test_join_across_three_files(new_style_files):
    rows = list(cermat_jpz.parse(new_style_files))
    r = next(r for r in rows if r["izo"] == "000638510" and r["zamereni_oboru"] == "")
    # z kapacity.xlsx / prihlasky.xlsx
    assert r["kapacita"] == 30
    assert r["prihlasky_celkem"] == 282
    # z vysledky.xlsx
    assert r["prijati"] == 30
    assert r["skor_prumer_cjma"] == 61.7
    assert r["neprijati_nesplneni_podminek"] == 2


def test_join_handles_is_so_rename(new_style_files, old_style_prihlasky_kapacity):
    """vysledky.xlsx má vždy ID_SO, ale prihlasky/kapacity 2024-2025 mají IS_SO
    - join přes ID_SOF musí fungovat bez ohledu na název sloupce ID_SO/IS_SO."""
    paths = {"vysledky": new_style_files["vysledky"], **old_style_prihlasky_kapacity}
    rows = list(cermat_jpz.parse(paths))
    r = next(r for r in rows if r["id_sof"] == "sof-1")
    assert r["kapacita"] == 30
    assert r["prihlasky_celkem"] == 282
    assert r["id_so"] == "so-1"  # zachyceno i když je ve zdrojovém souboru jako IS_SO


def test_file_with_all_blank_id_sof_warns_but_vysledky_still_wins(new_style_files, tmp_path, caplog):
    """Ověřeno živě na PZ2025_kolo2_skolobory_prihlasky.xlsx: CERMAT občas vydá
    soubor, kde je ID_SOF prázdné úplně ve všech řádcích (vlastní datová chyba
    zdroje). Neškodné, dokud vysledky.xlsx (nadmnožina sloupců) nese stejné
    metriky, ale musí se to aspoň zalogovat, ne jen tiše ignorovat."""
    rok, kolo = 2026, 1
    rows = [_vysledky_row("sof-1", "so-1", rok, kolo)]
    prihlasky_bez_id_sof = tmp_path / "prihlasky_bez_id_sof.xlsx"
    _build(prihlasky_bez_id_sof, rows, PRIHLASKY_EXTRA)

    wb = openpyxl.load_workbook(prihlasky_bez_id_sof)
    ws = wb.active
    header = [c.value for c in ws[1]]
    id_sof_col = header.index("ID_SOF") + 1
    for r in range(2, ws.max_row + 1):
        # ws.cell(value=None) je no-op v openpyxl (None = "hodnota nezadána"),
        # nutno smazat přímo přes .value.
        ws.cell(row=r, column=id_sof_col).value = None
    wb.save(prihlasky_bez_id_sof)

    paths = {"vysledky": new_style_files["vysledky"], "prihlasky": prihlasky_bez_id_sof}
    with caplog.at_level("WARNING"):
        rows_out = list(cermat_jpz.parse(paths))
    r = next(r for r in rows_out if r["id_sof"] == "sof-1")
    assert r["kapacita"] == 30  # z vysledky.xlsx, i když prihlasky.xlsx nepřispělo ničím
    assert r["prihlasky_celkem"] == 282
    assert "prázdné ID_SOF" in caplog.text


def test_vysledky_only_still_works(new_style_files):
    """Kdyby prihlasky/kapacita chyběly (např. 404), vysledky.xlsx samo o
    sobě obsahuje všechny požadované sloupce (je jejich nadmnožinou)."""
    rows = list(cermat_jpz.parse({"vysledky": new_style_files["vysledky"]}))
    r = next(r for r in rows if r["izo"] == "000638510" and r["zamereni_oboru"] == "")
    assert r["kapacita"] == 30
    assert r["prihlasky_celkem"] == 282
    assert r["prijati"] == 30


def test_import_and_idempotency(new_style_files):
    conn = db.connect(":memory:")
    rows = list(cermat_jpz.parse(new_style_files))
    stats1 = cermat_jpz.import_rows(conn, rows, url="http://x", soubor=new_style_files["vysledky"], rok=2026, kolo=1)
    assert stats1["prijimaci_rizeni"] == 3
    n1 = conn.execute("SELECT COUNT(*) FROM prijimaci_rizeni").fetchone()[0]

    stats2 = cermat_jpz.import_rows(conn, rows, url="http://x", soubor=new_style_files["vysledky"], rok=2026, kolo=1)
    n2 = conn.execute("SELECT COUNT(*) FROM prijimaci_rizeni").fetchone()[0]
    assert n1 == n2 == 3
    assert conn.execute("SELECT COUNT(*) FROM import_run WHERE zdroj = 'cermat_jpz'").fetchone()[0] == 2

    r = conn.execute(
        "SELECT kapacita, prihlasky_celkem, prijati FROM prijimaci_rizeni"
        " WHERE izo = '000638510' AND kod_kkov = ? AND zamereni_oboru = ''",
        (REF_KKOV,),
    ).fetchone()
    assert tuple(r) == (30, 282, 30)


def test_jen_praha_filter(new_style_files):
    conn = db.connect(":memory:")
    rows = list(cermat_jpz.parse(new_style_files))
    cermat_jpz.import_rows(conn, rows, jen_izo={"000638510"})
    izos = {r[0] for r in conn.execute("SELECT DISTINCT izo FROM prijimaci_rizeni")}
    assert izos == {"000638510"}


def test_download_returns_none_on_404(monkeypatch, tmp_path):
    class FakeResp:
        status_code = 404

    class FakeSession:
        @staticmethod
        def get(url, headers=None, timeout=None):
            return FakeResp()

    import requests
    monkeypatch.setattr(requests, "get", FakeSession.get)
    result = cermat_jpz.download(2099, 2, "vysledky", raw_dir=tmp_path)
    assert result is None


def test_parse_roky_range():
    assert cermat_jpz._parse_roky("2024-2026") == [2024, 2025, 2026]
    assert cermat_jpz._parse_roky("2026") == [2026]
    assert cermat_jpz._parse_roky("2024,2026") == [2024, 2026]
